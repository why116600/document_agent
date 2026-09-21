from pydantic import BaseModel, Field
from docx import Document
from docx.oxml.ns import qn
import json

try:
    from document_agent.agent.llm_client import llm_invoke, llm_model_invoke, extract_json_from_text
    from document_agent.agent.agent_core import AgentState
except ImportError:
    from llm_client import llm_invoke, llm_model_invoke, extract_json_from_text
    from agent_core import AgentState
from document_agent.write_tool.word_tool import replace_node_with_data,ParagraphItem,TableItem,DocxRoot,resolve_save_path,save_document

def create_new_docx_node(llm):#从零编写文档的节点
    def docx_node(state : AgentState) -> AgentState:#写docx的节点
        s=state["messages"]
        # input_files=state['input_file_path']
        # file_list_str="\n".join(input_files)
        retrieved_items=state["retrieved_content"]
        retrieved_info=""
        if retrieved_items is not None and len(retrieved_items)>0:
            retrieved_info="目前已知：\n"+"\n".join(retrieved_items)
        prompt=(
            "你是文档写作助手，现在需要用户输入和已知信息列出整个文档的各个章节\n"
            "如果某些内容可以按表格的形式输出，则按一个章节来输出\n"
            f"用户输入：{s}\n"
            f"{retrieved_info}\n"
            "输出结果按照json的列表格式输出章节列表，仅输出json字符串结果，不要输出其它内容"
        )
        res=llm_invoke(llm,prompt)
        section_str=getattr(res,"content")
        if section_str is None:
            return {**state,"state":"error","error":"规划文档章节时，模型输出错误"}
        matches = extract_json_from_text(section_str, return_all=True) or []
        print("分解结果：",matches)
        if len(matches)<=0:
            print("章节内容输出不包含json格式的数据")
        try:
            sections=json.loads(matches[0])
        except Exception as e:
            print("解析分段json数据失败：",e)
            return {**state,"state":"error","error":f"{e}"}
            
        doc=Document()
        style = doc.styles['Normal']
        style.font.name = 'Times New Roman'
        style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        for section in sections:
            print(f"编写{section}章节")
            prompt=(
                "你是一个写作系统的智能助手，需要根据章节主旨以及源信息进行文档生成，包括段落和表格\n"
                "文档字体上，除了用户的特殊要求，正文不加粗，标题加粗\n"
                f"整个文档的大纲：\n{section_str}\n"
                f"当前要编写的章节内容：{section}\n"
                f"{retrieved_info}\n"
            )
            res=llm_model_invoke(llm,prompt,DocxRoot)
            if res is None:
                print("写作失败")
                return {**state,"error":"写作失败","state":"error"}
            for item in res.items:
                if item.type=="paragraph":
                    node=doc.add_paragraph("")
                    replace_node_with_data(node,item.Paragraph)
                elif item.type=="table":
                    node=doc.add_table(rows=0,cols=0)
                    replace_node_with_data(node,item.Table)
        save_path=resolve_save_path(state.get("save_path"))
        print("保存到：",str(save_path))
        save_document(doc, save_path)
            
        # if action is None:
        #     print("识别用户写作意图失败")
        #     return {**state,"error":"识别用户写作意图失败","state":"error"}
        # if action.action_type=="new":
        #     prompt=(
        #         "你是一个写作系统的智能助手，需要根据用户的要求以及源信息进行文档生成，包括段落和表格\n"
        #         "文档字体上，除了用户的特殊要求，正文不加粗，标题加粗\n"
        #         f"用户的要求：{s}\n"
        #         f"{retrieved_info}\n"
        #     )
        #     res=llm_model_invoke(llm,prompt,DocxRoot)
        #     if res is None:
        #         print("写作失败")
        #         return {**state,"error":"写作失败","state":"error"}
        #     doc = Document()
        #     for item in res.items:
        #         if item.type=="paragraph":
        #             node=doc.add_paragraph("")
        #             replace_node_with_data(node,item.Paragraph)
        #         elif item.type=="table":
        #             node=doc.add_table(rows=0,cols=0)
        #             replace_node_with_data(node,item.Table)
        #     doc.save(action.save_path)
        return {**state,"state":"succeeded","save_path":str(save_path)}
    return docx_node
            