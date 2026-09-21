import copy
import os
import tempfile
import time
from pathlib import Path
from typing import Union, Optional, List
from pydantic import BaseModel, Field, model_validator
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from docx.table import Table, _Cell
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import RGBColor, Pt
from lxml import etree
import officemath2latex
import math2docx

#未指定格式时写回文档所使用的默认值
DEFAULT_FONT_SIZE = 12
DEFAULT_FONT_COLOR = "#000000"
DEFAULT_ALIGNMENT = "left"
DEFAULT_SPACING = 0
DEFAULT_LINE_SPACING = 1.0

#复制单元格格式时只搬运这些属性；合并信息（w:gridSpan/w:vMerge）保留新表格自己的
CELL_FORMAT_TAG_NAMES = frozenset([qn("w:tcBorders"), qn("w:shd"), qn("w:tcMar"), qn("w:vAlign")])

#段落/单元格内"模型无法表达、一旦整段清空就会永久丢失"的内容：图片、图形、文本框、域代码、公式对象等
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
PROTECTED_PARA_TAGS = (
    "w:drawing",
    "w:pict",
    "w:object",
    "w:fldSimple",
    "w:fldChar",
    "w:instrText",
    "w:txbxContent",
    "m:oMath",
    "m:oMathPara",
    f"{{{MC_NS}}}AlternateContent",
)


def _tag_name(tag: str) -> str:
    #已用Clark记法（{命名空间}标签）的直接使用，其余按docx前缀展开
    return tag if tag.startswith("{") else qn(tag)

class TextRunItem(BaseModel):
    text: str = Field(default="", description="文本内容")
    text_type: Optional[str] = Field(default="text", description="文本类型，text表示普通文本，latex表示latex格式的公式")
    bold: Optional[bool] = Field(default=None, description="是否加粗，None表示未指定，改写时沿用原文格式")
    italic: Optional[bool] = Field(default=None, description="是否斜体，None表示未指定，改写时沿用原文格式")
    underline: Optional[bool] = Field(default=None, description="是否下划线，None表示未指定，改写时沿用原文格式")
    font_size: Optional[int] = Field(default=None, description="字体大小，单位为磅，None表示未指定，改写时沿用原文格式")
    font_color: Optional[str] = Field(default=None, description="字体颜色，使用十六进制颜色代码，None表示未指定，改写时沿用原文格式")

    @model_validator(mode="before")
    @classmethod
    def _coerce_run(cls, data):
        if isinstance(data, str):
            return {"text": data, "text_type": "text"}
        if isinstance(data, dict):
            if "text_type" not in data or data["text_type"] is None:
                data["text_type"] = "text"
            if "text" not in data:
                data["text"] = ""
        return data

class ParagraphItem(BaseModel):
    runs: List[TextRunItem] = Field(default_factory=list, description="文本运行列表")
    alignment: Optional[str] = Field(default=None, description="段落对齐方式，可选值：left, center, right, justify，None表示未指定，改写时沿用原文格式")
    spacing_before: Optional[int] = Field(default=None, description="段前间距，单位为磅，None表示未指定，改写时沿用原文格式")
    spacing_after: Optional[int] = Field(default=None, description="段后间距，单位为磅，None表示未指定，改写时沿用原文格式")
    line_spacing: Optional[float] = Field(default=None, description="行间距倍数，例如 1.0 表示单倍行距，None表示未指定，改写时沿用原文格式")

    @model_validator(mode="before")
    @classmethod
    def _coerce_paragraph(cls, data):
        if isinstance(data, str):
            return {"runs": [{"text": data, "text_type": "text"}]}
        if isinstance(data, dict):
            if "runs" not in data and "text" in data:
                return {
                    "runs": [{"text": str(data["text"]), "text_type": "text", "bold": data.get("bold")}],
                    "alignment": data.get("alignment"),
                }
            if "runs" in data and isinstance(data["runs"], list):
                coerced_runs = []
                for r in data["runs"]:
                    if isinstance(r, str):
                        coerced_runs.append({"text": r, "text_type": "text"})
                    elif isinstance(r, dict):
                        if "text_type" not in r or r["text_type"] is None:
                            r["text_type"] = "text"
                        coerced_runs.append(r)
                    else:
                        coerced_runs.append(r)
                data["runs"] = coerced_runs
        return data

class GridItem(BaseModel):
    content: List[ParagraphItem] = Field(default_factory=list, description="单元格内容，包含一个或多个段落")
    row : int = Field(default=0, description="单元格所在的行索引，从0开始")
    col : int = Field(default=0, description="单元格所在的列索引，从0开始")
    row_span: int = Field(default=1, description="单元格跨越的行数")
    col_span: int = Field(default=1, description="单元格跨越的列数")

class TableItem(BaseModel):
    rows: int = Field(default=0, description="表格的行数")
    cols: int = Field(default=0, description="表格的列数")
    grid: List[GridItem] = Field(default_factory=list, description="表格的网格内容，包含每个单元格的内容和位置")

    @model_validator(mode="before")
    @classmethod
    def _coerce_table(cls, data):
        if isinstance(data, dict):
            if "cells" in data and isinstance(data["cells"], list) and not data.get("grid"):
                cells = data["cells"]
                rows = len(cells)
                cols = len(cells[0]) if rows > 0 and isinstance(cells[0], list) else 0
                grid = []
                for r_idx, row in enumerate(cells):
                    if isinstance(row, list):
                        for c_idx, cell_text in enumerate(row):
                            grid.append({
                                "row": r_idx,
                                "col": c_idx,
                                "row_span": 1,
                                "col_span": 1,
                                "content": [{"runs": [{"text": str(cell_text), "text_type": "text"}]}],
                            })
                return {"rows": rows, "cols": cols, "grid": grid}
        return data
    
class DocxItem(BaseModel):
    type: str = Field(description="文档元素类型，paragraph表示段落文本，table表示表格")
    Paragraph: Optional[ParagraphItem] = Field(default=None, description="段落文本内容")
    Table: Optional[TableItem] = Field(default=None, description="表格内容")
    
class DocxRoot(BaseModel):
    items: List[DocxItem] = Field(description="文档内容列表，每个元素表示一段文本或一个表格")
    
# ---------- 转换辅助函数 ----------

def _rgb_to_hex(rgb: Union[RGBColor, tuple, str]) -> str:
    """将 RGBColor 或 (r, g, b) 元组转换为 '#RRGGBB' 格式的十六进制字符串"""
    if rgb is None:
        return "#000000"
    if isinstance(rgb, str):
        return rgb if rgb.startswith("#") else f"#{rgb}"
    try:
        # RGBColor 继承自 tuple: (r, g, b)
        return f"#{rgb[0]:02x}{rgb[1]:02x}{rgb[2]:02x}"
    except Exception:
        pass
    try:
        if hasattr(rgb, "rgb"):
            return f"#{rgb.rgb:06x}"
    except Exception:
        pass
    return "#000000"

def _get_alignment(para: Paragraph) -> str:
    """将 WD_ALIGN_PARAGRAPH 枚举映射为字符串"""
    align = para.paragraph_format.alignment
    if align == WD_ALIGN_PARAGRAPH.LEFT:
        return "left"
    elif align == WD_ALIGN_PARAGRAPH.CENTER:
        return "center"
    elif align == WD_ALIGN_PARAGRAPH.RIGHT:
        return "right"
    elif align == WD_ALIGN_PARAGRAPH.JUSTIFY:
        return "justify"
    else:
        return "left"  # 默认

def _get_pt_value(val) -> int:
    """从 python-docx 的长度对象（如 Pt）中提取整数值（磅）"""
    if val is None:
        return 0
    if isinstance(val, Pt):
        return int(val.pt)
    return int(val)  # 兜底

def hex_to_rgb(hex_color: str) -> RGBColor:
    """将 '#RRGGBB' 转为 RGBColor 对象"""
    hex_color = hex_color.lstrip('#')
    r, g, b = int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16)
    return RGBColor(r, g, b)

def alignment_str_to_enum(align_str: str) -> WD_ALIGN_PARAGRAPH:
    """将对齐字符串映射为 WD_ALIGN_PARAGRAPH 枚举"""
    mapping = {
        'left': WD_ALIGN_PARAGRAPH.LEFT,
        'center': WD_ALIGN_PARAGRAPH.CENTER,
        'right': WD_ALIGN_PARAGRAPH.RIGHT,
        'justify': WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    return mapping.get(align_str, WD_ALIGN_PARAGRAPH.LEFT)

def _safe_get_font_color(font) -> Optional[Union[RGBColor, str]]:
    """安全读取 Run 的字体颜色，避免主题色或未初始化属性抛出异常"""
    if not font:
        return None
    try:
        color = font.color
        if not color:
            return None
        return color.rgb
    except Exception:
        return None

def convert_run(run: Run, preserve_none: bool = False) -> TextRunItem:
    """将单个 Run 转换为 TextRunItem。

    preserve_none 为 True 时，未直接设置的格式保留为 None（表示"由样式/原文决定"）。
    改写已有文档时用它取得基准格式，避免把样式继承来的格式写成显式默认值而覆盖原样式。
    """
    run_rgb = _safe_get_font_color(run.font)
    if preserve_none:
        return TextRunItem(
            text=run.text,
            text_type="text",
            bold=run.font.bold,
            italic=run.font.italic,
            underline=run.font.underline,
            font_size=_get_pt_value(run.font.size) if run.font.size is not None else None,
            font_color=_rgb_to_hex(run_rgb) if run_rgb is not None else None,
        )
    return TextRunItem(
        text=run.text,
        text_type="text",
        bold=run.font.bold if run.font.bold is not None else False,
        italic=run.font.italic if run.font.italic is not None else False,
        underline=run.font.underline if run.font.underline is not None else False,
        font_size=_get_pt_value(run.font.size) if run.font.size else 12,
        font_color=_rgb_to_hex(run_rgb) if run_rgb is not None else "#000000"
    )

def convert_paragraph(para: Paragraph, preserve_none: bool = False) -> ParagraphItem:
    """将 Paragraph 转换为 ParagraphItem，保留其所有 Run 的格式。

    preserve_none 为 True 时，未直接设置的段落格式保留为 None，含义同 convert_run。
    """
    runs=[]
    for child in para._element.iterchildren():
        tag=child.tag
        if tag==qn("w:r"):
            run = Run(child, para)
            runs.append(convert_run(run, preserve_none))
        elif tag in (qn('m:oMath'), qn('m:oMathPara')):
            omml_str= etree.tostring(child, encoding='unicode',pretty_print=True)
            latex = officemath2latex.process_math_string(omml_str)
            runs.append(TextRunItem(text=latex,text_type="latex"))
    if preserve_none:
        para_format = para.paragraph_format
        return ParagraphItem(
            runs=runs,
            alignment=_get_alignment(para) if para_format.alignment is not None else None,
            spacing_before=_get_pt_value(para_format.space_before) if para_format.space_before is not None else None,
            spacing_after=_get_pt_value(para_format.space_after) if para_format.space_after is not None else None,
            line_spacing=para_format.line_spacing,
        )
    return ParagraphItem(
        runs=runs,
        alignment=_get_alignment(para),
        spacing_before=_get_pt_value(para.paragraph_format.space_before),
        spacing_after=_get_pt_value(para.paragraph_format.space_after),
        line_spacing=para.paragraph_format.line_spacing or 1.0  # 若为 None 则默认为 1.0
    )

def convert_cell(cell: _Cell, row_idx: int, col_idx: int) -> GridItem:
    """将单元格转换为 GridItem，包含其中的所有段落"""
    return GridItem(
        content=[convert_paragraph(p) for p in cell.paragraphs],
        row=row_idx,
        col=col_idx,
        row_span=1,  # 将在表格转换时修正
        col_span=1   # 将在表格转换时修正
    )

def convert_table(table: Table) -> TableItem:
    """
    将 Table 转换为 TableItem，正确处理合并单元格。
    遍历时只记录每个合并区域的左上角单元格，并设置正确的 row_span 和 col_span。
    """
    rows = len(table.rows)
    cols = len(table.rows[0].cells) if rows > 0 else 0
    grid_items = []

    # 创建一个二维标记数组，记录某个网格位置是否已被处理（合并起始或已跳过）
    processed = [[False] * cols for _ in range(rows)]

    for r in range(rows):
        for c in range(cols):
            if processed[r][c]:
                continue

            cell = table.cell(r, c)
            # 获取合并信息
            vMerge = cell._tc.vMerge
            grid_span = cell._tc.grid_span

            # 计算合并跨度
            col_span = grid_span if grid_span is not None else 1
            # 计算纵向合并的行数：从当前行开始，找到连续 vMerge='continue' 的个数
            row_span = 1
            if vMerge == 'restart':
                # 向下寻找直到遇到 vMerge 为 None 或 'restart'
                for dr in range(r + 1, rows):
                    next_cell = table.cell(dr, c)
                    if next_cell._tc.vMerge == 'continue':
                        row_span += 1
                    else:
                        break
            # 标记所有被此合并覆盖的网格位置为已处理
            for rr in range(r, r + row_span):
                for cc in range(c, c + col_span):
                    if rr < rows and cc < cols:
                        processed[rr][cc] = True

            # 创建 GridItem（若单元格为空段落，content 可能为空列表，仍保留）
            grid_item = convert_cell(cell, r, c)
            grid_item.row_span = row_span
            grid_item.col_span = col_span
            grid_items.append(grid_item)

    return TableItem(rows=rows, cols=cols, grid=grid_items)

# ---------- 统一入口函数 ----------

def convert_node(node, preserve_none: bool = False) -> Union[TextRunItem, ParagraphItem, TableItem]:
    """
    根据节点类型转换为对应的 Pydantic 模型对象。

    支持的类型：
        - Run          -> TextRunItem
        - Paragraph    -> ParagraphItem
        - Table        -> TableItem
    其它类型抛出 TypeError。
    preserve_none 为 True 时，未直接设置的格式保留为 None（见 convert_run）。
    """
    if isinstance(node, Run):
        return convert_run(node, preserve_none)
    elif isinstance(node, Paragraph):
        return convert_paragraph(node, preserve_none)
    elif isinstance(node, Table):
        return convert_table(node)
    else:
        raise TypeError(f"不支持的节点类型: {type(node)}")
    
    
def _apply_paragraph_format(paragraph_format, para_item: ParagraphItem, write_defaults: bool = True) -> None:
    """把 ParagraphItem 的段落格式写到 paragraph_format 上。

    write_defaults 为 True 时，未设置的格式写显式默认值（新建文档用）；
    为 False 时，未设置的格式不写，交由原有样式继承（改写已有文档用，避免覆盖原格式）。
    """
    if write_defaults:
        paragraph_format.alignment = alignment_str_to_enum(para_item.alignment or DEFAULT_ALIGNMENT)
        paragraph_format.space_before = Pt(para_item.spacing_before or DEFAULT_SPACING)
        paragraph_format.space_after = Pt(para_item.spacing_after or DEFAULT_SPACING)
        paragraph_format.line_spacing = para_item.line_spacing or DEFAULT_LINE_SPACING
        return
    if para_item.alignment is not None:
        paragraph_format.alignment = alignment_str_to_enum(para_item.alignment)
    if para_item.spacing_before is not None:
        paragraph_format.space_before = Pt(para_item.spacing_before)
    if para_item.spacing_after is not None:
        paragraph_format.space_after = Pt(para_item.spacing_after)
    if para_item.line_spacing is not None:
        paragraph_format.line_spacing = para_item.line_spacing


def _apply_run_format(run: Run, run_item: TextRunItem, write_defaults: bool = True) -> None:
    """把 TextRunItem 的字体格式写到 run 上，write_defaults 的含义同 _apply_paragraph_format。"""
    if write_defaults:
        run.font.bold = bool(run_item.bold)
        run.font.italic = bool(run_item.italic)
        run.font.underline = bool(run_item.underline)
        run.font.size = Pt(run_item.font_size or DEFAULT_FONT_SIZE)
        run.font.color.rgb = hex_to_rgb(run_item.font_color or DEFAULT_FONT_COLOR)
        return
    if run_item.bold is not None:
        run.font.bold = run_item.bold
    if run_item.italic is not None:
        run.font.italic = run_item.italic
    if run_item.underline is not None:
        run.font.underline = run_item.underline
    if run_item.font_size is not None:
        run.font.size = Pt(run_item.font_size)
    if run_item.font_color is not None:
        run.font.color.rgb = hex_to_rgb(run_item.font_color)


def fill_table_from_grid(table: Table, grid_items, write_defaults: bool = True):
    """
    根据 GridItem 列表填充表格，处理合并单元格。
    注意：此函数假设表格已创建好行列数，且 grid_items 只包含合并区域的左上角单元格。
    write_defaults 的含义同 replace_node_with_data。
    """
    # 先填充所有单元格的段落内容（不带合并）
    for grid_item in grid_items:
        row, col = grid_item.row, grid_item.col
        cell = table.cell(row, col)
        # 移除多余段落，复用首个段落以避免留下空白首行
        first_p = cell.paragraphs[0]
        first_p._element.clear_content()
        for extra_p in list(cell.paragraphs[1:]):
            cell._tc.remove(extra_p._element)

        if not grid_item.content:
            continue

        for i, para_item in enumerate(grid_item.content):
            p = first_p if i == 0 else cell.add_paragraph()
            # 设置段落格式
            _apply_paragraph_format(p.paragraph_format, para_item, write_defaults)
            # 添加 runs
            for run_item in para_item.runs:
                if run_item.text_type == "latex":
                    math2docx.add_math(p, run_item.text)
                    continue
                run = p.add_run(run_item.text)
                _apply_run_format(run, run_item, write_defaults)
    # 处理合并（必须在内容填充之后，因为合并会改变单元格引用）
    for grid_item in grid_items:
        row, col = grid_item.row, grid_item.col
        row_span, col_span = grid_item.row_span, grid_item.col_span
        if row_span > 1 or col_span > 1:
            start_cell = table.cell(row, col)
            # 计算结束单元格位置
            end_row = row + row_span - 1
            end_col = col + col_span - 1
            end_cell = table.cell(end_row, end_col)
            # 合并
            start_cell.merge(end_cell)
            
def _copy_table_format(old_table: Table, new_table: Table) -> None:
    """把原表格的格式复制到重建的新表格上，避免改写表格时丢失外观。

    TableItem 只描述行列数与单元格内容，用 doc.add_table 重建会丢掉表格样式、边框、
    列宽、底纹等设置。这里复制表格级属性（w:tblPr/w:tblGrid）以及单元格级的边框、底纹、
    内边距、垂直对齐；合并信息（w:gridSpan/w:vMerge）保留新表格自己的，避免破坏结构。
    """
    old_tbl = old_table._element
    new_tbl = new_table._element

    if not old_tbl.findall(qn("w:tr")):
        # 原表格没有行（新建文档时用0行0列占位的空表格），没有格式可复制
        return

    old_tbl_pr = old_tbl.find(qn("w:tblPr"))
    old_grid = old_tbl.find(qn("w:tblGrid"))
    new_grid = new_tbl.find(qn("w:tblGrid"))
    cols_match = (
        old_grid is not None
        and new_grid is not None
        and len(old_grid.findall(qn("w:gridCol"))) == len(new_grid.findall(qn("w:gridCol")))
    )

    # 表格级属性：样式、边框、内边距、对齐、宽度等
    if old_tbl_pr is not None:
        copied_pr = copy.deepcopy(old_tbl_pr)
        if not cols_match:
            # 列数变化时不能沿用原表格宽度，交给 Word 按新列数重新计算
            old_w = copied_pr.find(qn("w:tblW"))
            if old_w is not None:
                copied_pr.remove(old_w)
        new_tbl_pr = new_tbl.find(qn("w:tblPr"))
        if new_tbl_pr is not None:
            new_tbl.remove(new_tbl_pr)
        new_tbl.insert(0, copied_pr)

    # 列宽定义：只在列数一致时沿用，否则保留新表格自己的列宽
    if cols_match:
        new_tbl.replace(new_grid, copy.deepcopy(old_grid))

    # 单元格级格式：逐行逐列搬运，保留新表格的合并信息
    for old_tr, new_tr in zip(old_tbl.findall(qn("w:tr")), new_tbl.findall(qn("w:tr"))):
        for old_tc, new_tc in zip(old_tr.findall(qn("w:tc")), new_tr.findall(qn("w:tc"))):
            old_tc_pr = old_tc.find(qn("w:tcPr"))
            if old_tc_pr is None:
                continue
            new_tc_pr = new_tc.find(qn("w:tcPr"))
            if new_tc_pr is None:
                new_tc_pr = OxmlElement("w:tcPr")
                new_tc.insert(0, new_tc_pr)
            for child in old_tc_pr:
                if child.tag not in CELL_FORMAT_TAG_NAMES:
                    continue
                existing = new_tc_pr.find(child.tag)
                if existing is not None:
                    new_tc_pr.remove(existing)
                new_tc_pr.append(copy.deepcopy(child))


def element_has_protected_content(element) -> bool:
    """判断元素（段落/表格/单元格）子树内是否含图片、图形、文本框、域代码等受保护内容。

    这类内容无法用 ParagraphItem/TableItem 表达，整段清空重建会永久丢失，因此需要保护。
    """
    if element is None:
        return False
    for tag in PROTECTED_PARA_TAGS:
        try:
            if element.findall(".//" + _tag_name(tag)):
                return True
        except Exception:
            continue
    return False


def paragraph_has_protected_content(para) -> bool:
    """判断段落内是否包含图片、图形、域代码等模型无法表达的内容。"""
    if not isinstance(para, Paragraph):
        return False
    return element_has_protected_content(para._element)


def _replace_text_runs_keep_protected(node: Paragraph, new_data: ParagraphItem,
                                      write_defaults: bool) -> None:
    """只重建段落的纯文本 run，保留图片等受保护元素（用于含图片段落的改写）。"""
    element = node._element
    has_new_math = any(r.text_type == "latex" for r in new_data.runs)
    for child in list(element):
        if child.tag == qn("w:pPr"):
            continue
        if child.tag == qn("w:hyperlink"):
            #超链接文字整体移除，避免与新增文本重复
            element.remove(child)
            continue
        if child.tag in (qn("m:oMath"), qn("m:oMathPara")):
            # 如果新内容显式提供了新的公式，则移除旧公式；否则原样保留
            if has_new_math:
                element.remove(child)
            continue
        if child.tag != qn("w:r"):
            continue
        if any(child.findall(".//" + _tag_name(tag)) for tag in PROTECTED_PARA_TAGS):
            continue  #含图片/域代码的 run 原样保留
        element.remove(child)
    _apply_paragraph_format(node.paragraph_format, new_data, write_defaults)
    for run_item in new_data.runs:
        if run_item.text_type == "latex":
            print("生成公式：", run_item.text)
            math2docx.add_math(node, run_item.text)
            continue
        run = node.add_run(run_item.text)
        _apply_run_format(run, run_item, write_defaults)


def replace_node_with_data(node, new_data, doc=None, write_defaults: bool = True):
    """
    原地替换文档中的节点（Paragraph 或 Table）为新数据定义的内容。

    参数：
        node: 当前节点（Paragraph 或 Table 对象）
        new_data: ParagraphItem 或 TableItem 对象
        doc: Document 对象（可选，若 node 未包含 part 属性则需传入）
        write_defaults: 未设置的格式是否写显式默认值。新建文档用 True（沿用原有行为）；
            改写已有文档用 False，让未设置的格式继续由原样式决定，避免把标题等样式覆盖掉。
    返回：
        替换后的新节点（Paragraph 或 Table 对象，对于段落仍返回原 node）
    """
    if doc is None:
        # 尝试从 node 获取文档
        if hasattr(node, 'part'):
            doc = node.part.document
        else:
            raise ValueError("无法获取 Document 对象，请传入 doc 参数")

    # ----- 替换段落 -----
    if isinstance(node, Paragraph) and isinstance(new_data, ParagraphItem):
        if paragraph_has_protected_content(node):
            #段内含图片/图形/域代码：整段清空会丢内容，改为只替换纯文本 run
            _replace_text_runs_keep_protected(node, new_data, write_defaults)
            return node
        # 清空所有子元素（保留段落本身与 w:pPr，即保留段落样式等原有格式）
        node._element.clear_content()
        # 设置段落格式
        _apply_paragraph_format(node.paragraph_format, new_data, write_defaults)
        # 添加新的 runs
        for run_item in new_data.runs:
            if run_item.text_type=="latex":
                print("生成公式：",run_item.text)
                math2docx.add_math(node, run_item.text)
                continue
            run = node.add_run(run_item.text)
            _apply_run_format(run, run_item, write_defaults)
        return node  # 返回原段落对象（已更新）

    # ----- 替换表格 -----
    elif isinstance(node, Table) and isinstance(new_data, TableItem):
        # 1. 在文档末尾创建新表格（填充内容）
        old_table = node
        new_table = doc.add_table(rows=new_data.rows, cols=new_data.cols)
        fill_table_from_grid(new_table, new_data.grid, write_defaults)

        # 2. 复制原表格的样式、边框、列宽、底纹等格式，避免改写后表格外观丢失
        _copy_table_format(old_table, new_table)

        # 3. 将新表格的元素移动到旧表格的位置
        old_element = old_table._element
        parent = old_element.getparent()
        index = parent.index(old_element)
        parent.insert(index, new_table._element)

        # 4. 删除旧表格元素
        parent.remove(old_element)

        # 5. 返回新表格对象（原对象已随旧元素失效），调用方需用它替换 nodes 中的引用
        return new_table

    else:
        raise TypeError(f"节点类型 {type(node)} 与数据 {type(new_data)} 不匹配")
    
# ---------- 保存相关 ----------

def resolve_save_path(save_path: Optional[str] = None, data_dir: str = "data",
                      name_template: str = "{index}", stem: str = "doc") -> Path:
    """
    计算文档的保存路径。

    参数：
        save_path: 用户指定的保存路径，为空时在 data_dir 目录下自动生成
        data_dir: 自动生成时使用的输出目录
        name_template: 自动生成时使用的文件名模板，可使用 {index} 和 {stem} 占位符
        stem: 文件名模板中 {stem} 的取值
    返回：
        保存路径（Path 对象），目录已创建
    """
    if save_path:
        path = Path(save_path)
        if path.suffix.lower() != ".docx":
            path = path.with_suffix(".docx")
        path.parent.mkdir(parents=True, exist_ok=True)
        return path
    dir_path = Path(data_dir)
    dir_path.mkdir(parents=True, exist_ok=True)
    index = 1
    candidate = dir_path / (name_template.format(index=index, stem=stem) + ".docx")
    while candidate.exists():
        index += 1
        candidate = dir_path / (name_template.format(index=index, stem=stem) + ".docx")
    return candidate


def _atomic_replace(tmp_path: Path, path: Path, retries: int = 5, delay: float = 0.05) -> None:
    """把临时文件原子替换为目标文件。

    Windows 上多个会话并发替换同一目标文件时，os.replace 可能瞬时返回"拒绝访问"，
    这里做有限次重试；重试用尽仍然失败就抛出，避免把真实错误吞掉。
    """
    for attempt in range(retries):
        try:
            os.replace(tmp_path, path)
            return
        except PermissionError:
            if attempt == retries - 1:
                raise
            time.sleep(delay)


def save_document(doc: Document, save_path) -> Path:
    """保存文档到指定路径，自动创建不存在的目录。

    先写到同目录下的唯一临时文件，再原子替换目标文件：
    这样即使保存过程中失败，目标文件也只会是原来的完整文件，不会留下半截损坏的文档；
    临时文件名唯一，多个会话同时保存同一输出路径时也不会互相覆盖。
    """
    path = Path(save_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    # 用 mkstemp 在目标目录下生成唯一临时文件，保证 os.replace 是同一文件系统内的原子替换
    fd, tmp_name = tempfile.mkstemp(dir=str(path.parent), prefix=path.name + ".", suffix=".tmp")
    os.close(fd)
    tmp_path = Path(tmp_name)
    try:
        doc.save(str(tmp_path))
        _atomic_replace(tmp_path, path)
    except Exception:
        # 保存失败时清理临时文件，避免在输出目录留下垃圾
        try:
            if tmp_path.exists():
                tmp_path.unlink()
        except OSError:
            pass
        raise
    return path