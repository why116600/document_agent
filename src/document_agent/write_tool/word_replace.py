import re
from typing import Dict, List, Tuple
from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table


def replace_in_paragraph(p: Paragraph, old_text: str, new_text: str) -> int:
    """在单个段落中安全替换文本，支持跨 Run（多个文本分块）的文本替换。

    返回该段落中成功替换的次数。
    """
    if not old_text or old_text not in p.text:
        return 0

    count = 0
    # 循环查找替换，直到该段落中不再包含 old_text
    while old_text in p.text:
        # 1. 尝试在单个 Run 内部替换，100% 保持该 Run 原格式
        single_run_replaced = False
        for run in p.runs:
            if old_text in run.text:
                run.text = run.text.replace(old_text, new_text, 1)
                count += 1
                single_run_replaced = True
                break

        if single_run_replaced:
            continue

        # 2. 如果 old_text 跨越了多个 Run，进行跨 Run 映射替换
        # 构建字符索引到 (run_idx, offset_in_run) 的映射
        char_map: List[Tuple[int, int]] = []
        for r_idx, run in enumerate(p.runs):
            for c_idx in range(len(run.text)):
                char_map.append((r_idx, c_idx))

        full_text = p.text
        start_pos = full_text.find(old_text)
        if start_pos < 0:
            break

        end_pos = start_pos + len(old_text)
        start_run_idx, start_char_idx = char_map[start_pos]
        end_run_idx, end_char_idx = char_map[end_pos - 1]

        # 跨 Run 拼接处理：
        # - start_run: 保留其前面部分，并追加 new_text
        # - 中间 runs: 清空其文本
        # - end_run: 保留其匹配终点后面的文本
        start_run = p.runs[start_run_idx]
        end_run = p.runs[end_run_idx]

        prefix = start_run.text[:start_char_idx]
        suffix = end_run.text[end_char_idx + 1:]

        if start_run_idx == end_run_idx:
            start_run.text = prefix + new_text + suffix
        else:
            start_run.text = prefix + new_text
            # 清空中间的 Run 文本
            for mid_idx in range(start_run_idx + 1, end_run_idx):
                p.runs[mid_idx].text = ""
            end_run.text = suffix

        count += 1

    return count


def replace_in_table(table: Table, old_text: str, new_text: str) -> int:
    """在表格的所有单元格段落中安全替换文本。"""
    count = 0
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                count += replace_in_paragraph(p, old_text, new_text)
    return count


def replace_in_document(doc: Document, old_text: str, new_text: str) -> int:
    """在整篇 Word 文档（正文所有段落 + 表格所有单元格）中查找替换文本。"""
    if not old_text:
        return 0

    total_count = 0
    # 1. 替换正文段落
    for p in doc.paragraphs:
        total_count += replace_in_paragraph(p, old_text, new_text)

    # 2. 替换表格单元格中的段落
    for table in doc.tables:
        total_count += replace_in_table(table, old_text, new_text)

    # 3. 替换页眉页脚（如有）
    for section in doc.sections:
        if section.header:
            for p in section.header.paragraphs:
                total_count += replace_in_paragraph(p, old_text, new_text)
            for t in section.header.tables:
                total_count += replace_in_table(t, old_text, new_text)
        if section.footer:
            for p in section.footer.paragraphs:
                total_count += replace_in_paragraph(p, old_text, new_text)
            for t in section.footer.tables:
                total_count += replace_in_table(t, old_text, new_text)

    return total_count


def batch_replace_in_document(doc: Document, replace_pairs: Dict[str, str]) -> Dict[str, int]:
    """批量替换文档中的多组键值对，返回每个目标词的替换命中次数统计。"""
    stats: Dict[str, int] = {}
    for old_text, new_text in replace_pairs.items():
        if not old_text:
            continue
        c = replace_in_document(doc, old_text, str(new_text))
        stats[old_text] = c
    return stats
